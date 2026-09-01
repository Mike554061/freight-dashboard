#!/usr/bin/env python3
"""
FleetView — Bid Vault document intake
=====================================

Pull clean, line-structured text out of a solicitation PDF or Word file so it
can go straight into the Bid Vault paste box.

    python3 tools/extract_questions.py rfp.pdf --clipboard
    python3 tools/extract_questions.py questionnaire.docx -o questions.txt
    python3 tools/extract_questions.py rfp.pdf --questions      # preview the split

WHY THIS IS A SEPARATE SCRIPT
-----------------------------
The browser cannot read a PDF, and adding a PDF library to a static page means
shipping a megabyte of parser to every page load. This runs locally, prints
text, and you paste it in. Nothing is uploaded anywhere.

WHAT IT DOES *NOT* DO
---------------------
It does not decide what a question is. Bid Vault's parseQuestions() in
js/bidvault-data.js is the single authority on that, so the two can never drift.
This script's job is to hand that parser text that still has its line structure:
numbered items on their own lines, form labels intact, wrapped prose rejoined.
--questions only PREVIEWS the split so you can sanity-check before pasting.

DEPENDENCIES
------------
PDF   : pymupdf (preferred, better layout) or pypdf — one of them, both common.
DOCX  : none. Reads the .docx zip directly, so table cells survive. Falls back
        to python-docx, then to macOS `textutil`.
"""

import argparse
import os
import re
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


# --------------------------------------------------------------------------
# Extraction — PDF
# --------------------------------------------------------------------------
def pdf_text(path):
    """Page text, best available extractor. Returns (text, engine)."""
    try:
        import fitz  # pymupdf
        doc = fitz.open(path)
        pages = [p.get_text("text") for p in doc]
        doc.close()
        return "\n\n".join(pages), "pymupdf"
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
    except ImportError:
        sys.exit("No PDF library found. Install one:  python3 -m pip install pymupdf")
    reader = PdfReader(path)
    return "\n\n".join((p.extract_text() or "") for p in reader.pages), "pypdf"


def pdf_form_fields(path):
    """AcroForm field names from a fillable PDF — these ARE the questions."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return []
    try:
        fields = PdfReader(path).get_fields() or {}
    except Exception:
        return []
    out = []
    for key, val in fields.items():
        label = ''
        if isinstance(val, dict):
            # /TU is the human-readable tooltip; it beats the internal /T name.
            label = str(val.get('/TU') or val.get('/T') or key)
        else:
            label = str(key)
        label = re.sub(r'\s+', ' ', label).strip()
        if label and label not in out:
            out.append(label)
    return out


# --------------------------------------------------------------------------
# Extraction — DOCX
# --------------------------------------------------------------------------
def docx_text(path):
    """Paragraphs and table cells, in document order. Returns (text, engine)."""
    try:
        return _docx_native(path), "docx-zip"
    except Exception:
        pass
    try:
        import docx  # python-docx, if the user happens to have it
        d = docx.Document(path)
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(' | '.join(c.text.strip() for c in row.cells))
        return "\n".join(lines), "python-docx"
    except ImportError:
        pass
    try:
        out = subprocess.run(['textutil', '-convert', 'txt', '-stdout', path],
                             capture_output=True, text=True, check=True)
        return out.stdout, "textutil"
    except Exception as e:
        sys.exit(f"Could not read {path}: {e}")


def _docx_native(path):
    """Read word/document.xml straight out of the zip — no dependencies, and
    table cells stay separate instead of being flattened into one blob."""
    with zipfile.ZipFile(path) as z:
        xml = z.read('word/document.xml')
    body = ET.fromstring(xml).find(f'{W_NS}body')
    if body is None:
        raise ValueError('no document body')

    def para_text(p):
        return ''.join(t.text or '' for t in p.iter(f'{W_NS}t')).strip()

    lines = []
    for el in body:
        tag = el.tag
        if tag == f'{W_NS}p':
            lines.append(para_text(el))
        elif tag == f'{W_NS}tbl':
            for row in el.findall(f'{W_NS}tr'):
                raw_cells = [' '.join(para_text(p) for p in c.findall(f'{W_NS}p')).strip()
                             for c in row.findall(f'{W_NS}tc')]
                cells = [c for c in raw_cells if c]
                if not cells:
                    continue
                # A vendor-form row is "label | answer space". The answer space
                # is usually EMPTY, so keying off len(cells)==2 misses it — the
                # signal is that the row HAS more than one cell, not that both
                # are filled. Emitting "Label:" is what makes the Bid Vault
                # parser see a form field rather than a stray sentence.
                is_label_row = (len(raw_cells) >= 2 and len(cells) <= 2
                                and len(cells[0]) < 70)
                if is_label_row:
                    lines.append(cells[0].rstrip(':') + ':')
                else:
                    lines.append(' | '.join(cells))
    return "\n".join(lines)


# --------------------------------------------------------------------------
# Normalisation — keep the structure the browser parser relies on
# --------------------------------------------------------------------------
NUM_START = re.compile(r'^\s*(\d+[\.\)]|\(\d+\)|[a-zA-Z][\.\)]|[-•*·▪])\s+')
FIELD_END = re.compile(r'(:|_{3,})\s*$')
PAGE_ONLY = re.compile(r'^\s*(page\s+)?\d+(\s+of\s+\d+)?\s*$', re.I)
SENT_END  = re.compile(r'[.!?:;]\s*$')


HEADING = re.compile(r'^[^a-z]+$')          # no lowercase anywhere = section furniture


def is_heading(line):
    """A section heading, not a question. Deliberately narrow: an ALL-CAPS line
    that neither asks anything nor labels a field. 'LEGAL BUSINESS NAME:' keeps
    its colon and survives; 'SECTION L - INSTRUCTIONS TO OFFERORS' does not."""
    s = line.strip()
    if not s or len(s) > 80:
        return False
    if s.endswith(':') or s.endswith('?') or s.endswith('_'):
        return False
    if NUM_START.match(s):
        return False
    return bool(HEADING.match(s)) and any(c.isalpha() for c in s)


def normalise(text, keep_headings=False):
    # De-hyphenate words split across a line break.
    text = re.sub(r'(\w)-\n(\w)', r'\1\2', text)
    raw = [ln.rstrip() for ln in text.replace('\r\n', '\n').replace('\r', '\n').split('\n')]

    kept = []
    for ln in raw:
        s = ln.strip()
        if not s:
            kept.append('')
            continue
        if PAGE_ONLY.match(s):          # bare page numbers
            continue
        if len(s) <= 2 and not s.isdigit():
            continue
        if not keep_headings and is_heading(s):
            continue
        kept.append(s)

    # Rejoin prose that PDF extraction wrapped mid-sentence, WITHOUT welding
    # separate numbered items or form labels together.
    out = []
    for s in kept:
        if not s:
            out.append('')
            continue
        prev = out[-1] if out else ''
        starts_item = bool(NUM_START.match(s))
        is_label = bool(FIELD_END.search(s))
        prev_open = (prev and not SENT_END.search(prev) and not FIELD_END.search(prev)
                     and not prev.isupper())
        continues = (prev_open and not starts_item and not is_label
                     and (s[0].islower() or s[0] in '(,'))
        if continues:
            out[-1] = prev + ' ' + s
        else:
            out.append(s)

    text = '\n'.join(out)
    return re.sub(r'\n{3,}', '\n\n', text).strip()


# --------------------------------------------------------------------------
# Preview only — mirrors js/bidvault-data.js parseQuestions()
# --------------------------------------------------------------------------
def preview_questions(text):
    lines = text.split('\n')
    out, buf = [], ''

    def flush():
        nonlocal buf
        t = re.sub(r'\s+', ' ', buf.strip())
        clean = re.sub(r'^\s*(\d+[\.\)]|[a-zA-Z][\.\)]|[-•*·])\s*', '', t).strip()
        if len(clean) >= 3:
            out.append(clean)
        buf = ''

    for raw in lines:
        line = raw.strip()
        if not line:
            if buf:
                flush()
            continue
        if NUM_START.match(line) or FIELD_END.search(line):
            if buf:
                flush()
            buf = FIELD_END.sub('', line)
            flush()
            continue
        buf = (buf + ' ' + line) if buf else line
        if line.endswith('?'):
            flush()
    if buf:
        flush()
    return out


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(
        description="Extract solicitation text for the FleetView Bid Vault.")
    ap.add_argument('file', help='.pdf, .docx, .doc, .rtf, .txt, .md or .csv')
    ap.add_argument('-o', '--out', help='write to this file instead of stdout')
    ap.add_argument('--clipboard', action='store_true',
                    help='copy the result to the clipboard (macOS pbcopy)')
    ap.add_argument('--questions', action='store_true',
                    help='preview how Bid Vault will split it, one per line')
    ap.add_argument('--raw', action='store_true',
                    help='skip normalisation and dump extractor output as-is')
    ap.add_argument('--keep-headings', action='store_true',
                    help='keep ALL-CAPS section headings (dropped by default as furniture)')
    args = ap.parse_args()

    path = os.path.expanduser(args.file)
    if not os.path.isfile(path):
        sys.exit(f"No such file: {path}")
    ext = os.path.splitext(path)[1].lower()

    fields = []
    if ext == '.pdf':
        text, engine = pdf_text(path)
        fields = pdf_form_fields(path)
    elif ext in ('.docx', '.doc', '.rtf', '.odt'):
        text, engine = docx_text(path)
    elif ext in ('.txt', '.md', '.csv'):
        with open(path, encoding='utf-8', errors='replace') as fh:
            text = fh.read()
        engine = 'plain'
    else:
        sys.exit(f"Unsupported file type: {ext}")

    if not args.raw:
        text = normalise(text, keep_headings=args.keep_headings)

    # Fillable-PDF field names are the cleanest question source there is —
    # put them first, labelled, so the parser reads them as form fields.
    if fields:
        # The page text also PRINTS these labels next to their boxes, so without
        # this the same field arrives twice — once as a field, once as prose.
        seen = {f.rstrip(':').strip().lower() for f in fields}
        body_lines = [ln for ln in text.split('\n')
                      if ln.strip().rstrip(':').lower() not in seen]
        # No commentary lines in the output — everything printed here gets
        # pasted into Bid Vault, and a "# heading" would parse as a question.
        # The field count goes to stderr in the summary instead.
        text = ('\n'.join(f.rstrip(':') + ':' for f in fields) + '\n\n'
                + re.sub(r'\n{3,}', '\n\n', '\n'.join(body_lines)).strip()).strip()

    body = '\n'.join(preview_questions(text)) if args.questions else text

    if args.out:
        with open(args.out, 'w', encoding='utf-8') as fh:
            fh.write(body + '\n')
        print(f"Wrote {args.out}", file=sys.stderr)
    if args.clipboard:
        try:
            subprocess.run(['pbcopy'], input=body, text=True, check=True)
            print("Copied to clipboard — paste into Bid Vault.", file=sys.stderr)
        except Exception as e:
            print(f"Clipboard copy failed: {e}", file=sys.stderr)
    if not args.out and not args.clipboard:
        print(body)

    n = len(preview_questions(text))
    print(f"[{engine}] {len(text.split(chr(10)))} lines · "
          f"{len(fields)} form fields · ~{n} questions", file=sys.stderr)


if __name__ == '__main__':
    main()
